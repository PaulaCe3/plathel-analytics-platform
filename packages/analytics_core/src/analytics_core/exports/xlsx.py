"""Write-only XLSX workbook, buffered in memory and delivered in chunks."""
from datetime import datetime
from io import BytesIO
import json
from openpyxl import Workbook
from analytics_core.exports.security import safe_cell, unique_headers


def _cell(value):
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, datetime) and value.tzinfo is not None:
        value = value.isoformat()
    return safe_cell(value)


class XLSXExporter:
    format = "xlsx"
    content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def export(self, source, options):
        workbook = Workbook(write_only=True)
        output = BytesIO()
        try:
            data = workbook.create_sheet("Datos")
            data.append(list(unique_headers([_cell(value) for value in source.headers])))
            for row in source.rows():
                data.append([_cell(value) for value in row])
            if options.include_transformations:
                audit = workbook.create_sheet("Registro de cambios")
                audit.append(["Orden", "ID", "Acción", "Tipo", "Columnas", "Parámetros", "Filas antes", "Filas después", "Filas afectadas", "Descripción", "Ejemplos", "Automática", "Fecha"])
                for item in source.transformations.transformations:
                    audit.append([_cell(value) for value in [item.seq, item.id, item.action_id, item.kind, item.columns, item.params, item.rows_before, item.rows_after, item.rows_affected, item.summary, [example.model_dump() for example in item.examples], item.automatic, item.at]])
            summary = workbook.create_sheet("Resumen")
            summary.append(["Concepto", "Valor"])
            for key, value in source.summary.items():
                summary.append([_cell(key), _cell(value)])
            workbook.save(output)
        except Exception:
            # openpyxl uses temporary worksheet XML internally; always remove it.
            for sheet in workbook.worksheets:
                if sheet._writer:
                    if not sheet.closed:
                        sheet.close()
                    sheet._writer.cleanup()
            output.close()
            raise
        finally:
            workbook.close()
        output.seek(0)
        return self._chunks(output)

    @staticmethod
    def _chunks(output):
        try:
            while chunk := output.read(64 * 1024):
                yield chunk
        finally:
            output.close()
