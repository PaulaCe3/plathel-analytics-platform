"""Pandas implementation of the Phase 1 data engine."""

import csv
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from analytics_core.errors import AppError
from analytics_core.ingestion.headers import normalize_headers, stable_column_key
from analytics_core.ingestion.models import ColumnMetadata, ParseResult, SourceSettings
from analytics_core.ingestion.sniffing import sniff_csv
from analytics_core.canonical.derived import DerivedFieldRule
from analytics_core.canonical.fields import FieldSpec
from analytics_core.cleaning.models import CanonicalBuildResult, CleaningActionSpec, CleaningPlan
from analytics_core.engine.pandas_impl.checks import cleaning_plan, inspect
from analytics_core.engine.pandas_impl.parsers import build_frame, persist_result
from analytics_core.engine.pandas_impl.transforms import apply_actions
from analytics_core.mapping.models import ColumnMapping
from analytics_core.quality.models import DataQualityReport
from analytics_core.validation.models import ParseReport, ProfileCheck
from analytics_core.engine.query import DateCoverage, QueryResult, QuerySpec
from analytics_core.engine.pandas_impl.query import date_coverage as query_date_coverage
from analytics_core.engine.pandas_impl.query import run_query as execute_query


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return str(value.isoformat())
    return str(value)


def _header_score(row: tuple[Any, ...]) -> tuple[float, int]:
    values = [_stringify(value).strip() for value in row]
    nonempty = [value for value in values if value]
    if not nonempty:
        return (0.0, 0)
    unique = len(set(nonempty)) / len(nonempty)
    text = sum(not value.replace(".", "", 1).isdigit() for value in nonempty) / len(nonempty)
    return (unique + text, len(nonempty))


def _detect_header(rows: list[tuple[Any, ...]]) -> int:
    candidates = [(_header_score(row), index) for index, row in enumerate(rows[:25])]
    score, index = max(candidates, key=lambda item: item[0], default=((0.0, 0), 0))
    if score[1] == 0:
        raise AppError(code="HEADER_NOT_FOUND", http_status=422, message="No se encontró una fila de encabezados.")
    return index


def _metadata(frame: pd.DataFrame, original_names: list[str]) -> list[ColumnMetadata]:
    result: list[ColumnMetadata] = []
    total = len(frame.index)
    for index, column in enumerate(frame.columns, start=1):
        series = frame[column].fillna("").astype(str)
        nonempty = series[series.str.strip() != ""]
        samples = nonempty.head(5).tolist()
        result.append(
            ColumnMetadata(
                key=stable_column_key(index),
                original_name=original_names[index - 1],
                sample=samples,
                null_ratio=round(1 - (len(nonempty) / total), 4) if total else 0.0,
                approximate_cardinality=int(nonempty.nunique()) if len(nonempty) <= 100_000 else None,
            )
        )
    return result


class PandasDataEngine:
    def parse_to_parquet(
        self,
        source_path: Path,
        extension: str,
        destination: Path,
        settings: SourceSettings,
        *,
        max_rows: int,
        max_columns: int,
    ) -> ParseResult:
        if extension == ".csv":
            frame, originals, resolved, sheets = self._read_csv(source_path, settings, max_rows)
            selected = None
        else:
            frame, originals, resolved, sheets, selected = self._read_xlsx(source_path, settings, max_rows)
        if len(frame.columns) > max_columns:
            raise AppError(code="TOO_MANY_COLUMNS", http_status=413, message=f"El archivo supera el máximo de {max_columns} columnas.")
        if len(frame.index) > max_rows:
            raise AppError(code="TOO_MANY_ROWS", http_status=413, message=f"El archivo supera el máximo de {max_rows} filas.")
        frame = frame.fillna("").astype(str)
        frame.columns = [stable_column_key(index) for index in range(1, len(frame.columns) + 1)]
        destination.parent.mkdir(parents=True, exist_ok=True)
        frame.to_parquet(destination, index=False)
        return ParseResult(
            row_count=len(frame.index),
            column_count=len(frame.columns),
            columns=_metadata(frame, originals),
            source_settings=resolved,
            available_sheets=sheets,
            selected_sheet=selected,
        )

    def _read_csv(self, path: Path, settings: SourceSettings, max_rows: int):
        encoding, delimiter = sniff_csv(path, settings.encoding, settings.delimiter)
        header_row = settings.header_row or 0
        try:
            with path.open("r", encoding=encoding, newline="") as source:
                rows = csv.reader(source, delimiter=delimiter)
                header = next((row for index, row in enumerate(rows) if index == header_row), None)
            if not header:
                raise AppError(code="HEADER_NOT_FOUND", http_status=422, message="No se encontró una fila de encabezados.")
            originals, normalized = normalize_headers(header)
            frame = pd.read_csv(
                path,
                encoding=encoding,
                sep=delimiter,
                header=None,
                skiprows=header_row + 1,
                names=normalized,
                dtype=str,
                keep_default_na=False,
                na_filter=False,
                engine="python",
                quoting=csv.QUOTE_MINIMAL,
                nrows=max_rows + 1,
                skip_blank_lines=True,
            )
        except (UnicodeError, pd.errors.ParserError, pd.errors.EmptyDataError) as exc:
            raise AppError(code="FILE_CORRUPT", http_status=422, message="No se pudo leer el archivo CSV.") from exc
        resolved = settings.model_copy(update={"encoding": encoding, "delimiter": delimiter, "header_row": header_row})
        return frame, originals, resolved, []

    def _read_xlsx(self, path: Path, settings: SourceSettings, max_rows: int):
        source = path.open("rb")
        try:
            workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
        except Exception as exc:
            source.close()
            raise AppError(code="FILE_CORRUPT", http_status=422, message="El archivo XLSX está dañado o no es válido.") from exc
        try:
            sheets = []
            for worksheet in workbook.worksheets:
                if any(any(value is not None and str(value).strip() for value in row) for row in worksheet.iter_rows(min_row=1, max_row=25, values_only=True)):
                    sheets.append(worksheet.title)
            if not sheets:
                raise AppError(code="FILE_EMPTY", http_status=422, message="El archivo XLSX no contiene hojas con datos.")
            selected = settings.sheet or sheets[0]
            if selected not in sheets:
                raise AppError(code="SHEET_NOT_FOUND", http_status=404, message="La hoja solicitada no existe.")
            worksheet = workbook[selected]
            rows = list(worksheet.iter_rows(values_only=True))
        finally:
            workbook.close()
            source.close()
        header_row = settings.header_row if settings.header_row is not None else _detect_header(rows)
        if header_row >= len(rows):
            raise AppError(code="HEADER_NOT_FOUND", http_status=422, message="La fila de encabezados no existe.")
        originals, normalized = normalize_headers(rows[header_row])
        data_rows = rows[header_row + 1 : header_row + 2 + max_rows]
        width = len(normalized)
        values = [[_stringify(value) for value in tuple(row)[:width]] + [""] * max(0, width - len(tuple(row))) for row in data_rows]
        frame = pd.DataFrame(values, columns=normalized)
        frame = frame.loc[~frame.apply(lambda row: all(not str(value).strip() for value in row), axis=1)]
        resolved = settings.model_copy(update={"sheet": selected, "header_row": header_row})
        return frame, originals, resolved, sheets, selected

    def preview(self, parquet_path: Path, rows: int) -> list[dict[str, str | None]]:
        frame = pd.read_parquet(parquet_path).head(rows)
        return [{key: (None if value == "" else str(value)) for key, value in record.items()} for record in frame.to_dict(orient="records")]

    def build_canonical(
        self,
        raw_path: Path,
        destination: Path,
        mappings: list[ColumnMapping],
        fields: list[FieldSpec],
        column_names: dict[str, str],
        settings: SourceSettings,
        derived_rules: list[DerivedFieldRule],
        actions: list[CleaningActionSpec],
    ) -> CanonicalBuildResult:
        frame, reports, automatic = build_frame(raw_path, mappings, fields, column_names, settings, derived_rules)
        frame, selected = apply_actions(frame, actions, len(automatic))
        return persist_result(frame, destination, reports, [*automatic, *selected])

    def inspect_quality(self, canonical_path: Path, reports: list[ParseReport], profile_checks: list[ProfileCheck]) -> DataQualityReport:
        return inspect(canonical_path, reports, profile_checks)

    def create_cleaning_plan(self, canonical_path: Path, reports: list[ParseReport], quality: DataQualityReport) -> CleaningPlan:
        return cleaning_plan(canonical_path, reports, quality)

    def run_query(self, canonical_path: Path, query: QuerySpec) -> QueryResult:
        return execute_query(canonical_path, query)

    def date_coverage(self, canonical_path: Path, field: str) -> DateCoverage:
        return query_date_coverage(canonical_path, field)
