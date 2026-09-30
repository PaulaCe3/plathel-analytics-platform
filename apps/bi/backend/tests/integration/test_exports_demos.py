import csv
import io
import json
from datetime import UTC, datetime, timedelta
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from analytics_core.settings import Settings
from analytics_core.engine.pandas_impl import PandasDataEngine
from analytics_core.engine.query import FilterClause, MeasureSpec, QuerySpec
from analytics_core.sessions.store import DatasetSessionStore
from bi.api.main import create_app
from bi.demo_registry import DEMO_ROOT


@pytest.fixture
def app_client(tmp_path):
    settings = Settings(dataset_storage_path=tmp_path)
    with TestClient(create_app(settings), raise_server_exceptions=False) as client:
        yield client, settings


def own_dataset(client, data=None):
    csv_data = data or "fecha,producto,importe,provincia,canal,moneda,segmento,presupuesto,ignorada\n2026-01-01,Mate,100.5,Córdoba,Online,ARS,A,50,original uno\n2026-01-01,Mate,100.5,Córdoba,Online,ARS,A,50,original dos\n2026-02-01,Café,-42.5,Córdoba,Online,ARS,B,25,original tres\n2026-02-02,Termo,300,Buenos Aires,Local,ARS,A,60,original cuatro\n"
    uploaded = client.post("/api/v1/datasets", files={"file": ("cliente_privado.xlsx" if isinstance(csv_data, bytes) else "cliente_privado.csv", csv_data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if isinstance(csv_data, bytes) else "text/csv")})
    assert uploaded.status_code == 201, uploaded.text
    dataset = uploaded.json()["dataset_id"]
    fields = ["date", "concept", "amount", "location", "channel", "currency", None, None, None]
    dispositions = ["canonical"] * 6 + ["custom_dimension", "custom_measure", "ignored"]
    mappings = [{"column_key": f"c{i:02d}", "target_field": field, "disposition": disposition} for i, (field, disposition) in enumerate(zip(fields, dispositions), 1)]
    mapped = client.put(f"/api/v1/datasets/{dataset}/mapping", json={"profile_id": "retail_ecommerce", "mappings": mappings})
    assert mapped.status_code == 200, mapped.text
    assert client.post(f"/api/v1/datasets/{dataset}/validate").status_code == 200
    return dataset


def ready(client, dataset, actions=()):
    response = client.put(f"/api/v1/datasets/{dataset}/cleaning", json={"actions": list(actions)})
    assert response.status_code == 200, response.text
    return response.json()


def parse_csv(response, delimiter=";"):
    assert response.status_code == 200, response.text
    assert response.content.startswith(b"\xef\xbb\xbf")
    return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=delimiter))


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_own_file_full_pipeline_filtered_export_and_raw_join(app_client, format):
    client, settings = app_client
    dataset = own_dataset(client)
    raw_path = settings.dataset_storage_path / dataset / "raw.parquet"
    raw_before = raw_path.read_bytes()
    cleaned = ready(client, dataset, [{"id": "drop_exact_duplicates", "description": "Eliminar duplicados", "destructive": True}])
    assert cleaned["canonical_row_count"] == 3
    filters = [{"field": "location", "op": "in", "values": ["Córdoba"]}, {"field": "channel", "op": "in", "values": ["Online"]}, {"field": "date", "op": "between", "values": ["2026-02-01", "2026-02-28"]}]
    dashboard = client.post(f"/api/v1/datasets/{dataset}/dashboard", json={"filters": filters}).json()
    assert dashboard["filtered_row_count"] == 1
    before_files = set(path.name for path in raw_path.parent.iterdir())
    response = client.post(f"/api/v1/datasets/{dataset}/export", json={"format": format, "scope": "filtered_data", "filters": filters, "headers": "original", "include_original_columns": True, "include_ignored_columns": True})
    assert response.status_code == 200, response.text
    assert response.headers["content-disposition"] == f'attachment; filename="datos.{format}"'
    assert "cliente_privado" not in response.headers["content-disposition"]
    assert raw_before == raw_path.read_bytes()
    assert set(path.name for path in raw_path.parent.iterdir()) == before_files
    if format == "csv":
        rows = parse_csv(response)
        assert len(rows) == 2 and rows[1][2] == "-42,5"
        assert rows[1][-1] == "original tres"
    else:
        workbook = load_workbook(io.BytesIO(response.content), data_only=False)
        assert workbook.sheetnames == ["Datos", "Registro de cambios", "Resumen"]
        rows = list(workbook["Datos"].values)
        assert len(rows) == 2 and rows[1][2] == -42.5
        assert rows[1][-1] == "original tres"
        audit = list(workbook["Registro de cambios"].values)
        assert any(row[2] == "drop_exact_duplicates" and row[8] == 1 for row in audit[1:])
        summary = dict(list(workbook["Resumen"].values)[1:])
        assert summary["Filas exportadas"] == 1 and summary["Filas canonical"] == 3
        assert summary["Moneda"] == "ARS" and json.loads(summary["Filtros"]) == filters
        assert "cliente_privado" not in str(summary) and str(settings.dataset_storage_path) not in str(summary)
    assert len(set(rows[0])) == len(rows[0])
    assert "segmento" in rows[0] and "presupuesto" in rows[0]
    engine_count = PandasDataEngine().run_query(raw_path.parent / "canonical.parquet", QuerySpec(measures=[MeasureSpec(alias="rows", aggregation="count")], filters=[FilterClause(**item) for item in filters])).rows[0]["rows"]
    assert len(rows) - 1 == engine_count
    complete = client.post(f"/api/v1/datasets/{dataset}/export", json={"format": "csv", "scope": "clean_data", "filters": filters, "include_ignored_columns": True})
    complete_rows = parse_csv(complete)
    assert len(complete_rows) == 4
    assert all("original dos" not in row for row in complete_rows)
    assert complete_rows[0][1] == "Producto"


@pytest.mark.parametrize("format", ["csv", "xlsx"])
@pytest.mark.parametrize("payload", ["=1+1", "+cmd", "-cmd", "@SUM(A1)", "\tformula", "\rformula"])
def test_formula_injection_all_text_surfaces_and_negative_numbers(app_client, format, payload):
    client, _ = app_client
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow(["fecha", "producto", "importe", "provincia", "canal", "moneda", "segmento", "presupuesto", payload])
    writer.writerow(["2026-01-01", payload, -42.5, "Córdoba", "Online", "ARS", payload, -7.5, payload])
    source_data = buffer.getvalue()
    if payload.startswith("\r"):
        # XLSX preserves leading CR without exercising CSV dialect inference.
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["fecha", "producto", "importe", "provincia", "canal", "moneda", "segmento", "presupuesto", payload])
        sheet.append(["2026-01-01", payload, -42.5, "Córdoba", "Online", "ARS", payload, -7.5, payload])
        binary = io.BytesIO()
        workbook.save(binary)
        source_data = binary.getvalue()
    dataset = own_dataset(client, source_data)
    ready(client, dataset)
    response = client.post(f"/api/v1/datasets/{dataset}/export", json={"format": format, "headers": "original", "include_original_columns": True, "include_ignored_columns": True})
    assert response.status_code == 200, response.text
    if format == "csv":
        rows = parse_csv(response)
        assert rows[1][2] == "-42,5"
        assert rows[1][7] == "-7,5"
    else:
        workbook = load_workbook(io.BytesIO(response.content), data_only=False)
        rows = list(workbook["Datos"].values)
        assert rows[1][2] == -42.5 and rows[1][7] == -7.5
        assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    serialized_payload = payload.replace("\r", "\n") if isinstance(source_data, bytes) else payload
    assert rows[1][1] == "'" + serialized_payload
    assert rows[1][6] == "'" + serialized_payload
    assert rows[1][-1] == "'" + serialized_payload
    assert rows[0][-1] == "'" + serialized_payload


@pytest.mark.parametrize("demo_id,industry,metric,expected", [("retail_demo", "retail_ecommerce", "gross_profit", 120000), ("services_demo", "services", "service_hours", 36), ("hospitality_demo", "hospitality", "total_nights", 24)])
def test_demo_to_dashboard_to_export_uses_real_assets(app_client, demo_id, industry, metric, expected):
    client, settings = app_client
    listed = client.get("/api/v1/demos")
    assert listed.status_code == 200
    assert {item["id"] for item in listed.json()} == {"retail_demo", "services_demo", "hospitality_demo"}
    created = client.post("/api/v1/datasets/demo", json={"demo_id": demo_id})
    assert created.status_code == 201, created.text
    session = created.json()
    assert session["industry_id"] == industry and session["demo_id"] == demo_id and session["stage"] == "mapped"
    dataset = session["dataset_id"]
    metadata = json.loads((DEMO_ROOT / demo_id / "demo.json").read_text(encoding="utf-8"))
    current_mapping = client.get(f"/api/v1/datasets/{dataset}/mapping").json()["mappings"]
    assert current_mapping == metadata["preset_mapping"]
    validation = client.post(f"/api/v1/datasets/{dataset}/validate")
    assert validation.status_code == 200 and validation.json()["validation"]["valid"]
    ready(client, dataset)
    dashboard = client.post(f"/api/v1/datasets/{dataset}/dashboard", json={})
    assert dashboard.status_code == 200, dashboard.text
    body = dashboard.json()
    assert body["data"][f"kpi_{metric}"]["value"] == pytest.approx(expected)
    assert all(value.get("status") != "error" for value in body["data"].values())
    assert body["data"]["insights_top"]["insights"]
    terminology = {"retail_ecommerce": "Producto", "services": "Servicio", "hospitality": "Habitación"}
    assert body["spec"]["terminology"]["concept"] == terminology[industry]
    filters = [{"field": "channel", "op": "in", "values": ["Online"]}]
    filtered = client.post(f"/api/v1/datasets/{dataset}/dashboard", json={"filters": filters}).json()
    assert filtered["filtered_row_count"] == 8
    files_before = set((settings.dataset_storage_path / dataset).iterdir())
    for format in ("csv", "xlsx"):
        response = client.post(f"/api/v1/datasets/{dataset}/export", json={"format": format, "scope": "filtered_data", "filters": filters})
        assert response.status_code == 200, response.text
        if format == "csv":
            rows = parse_csv(response)
        else:
            workbook = load_workbook(io.BytesIO(response.content), data_only=False)
            rows = list(workbook["Datos"].values)
            assert dict(list(workbook["Resumen"].values)[1:])["Filas exportadas"] == 8
        assert len(rows) - 1 == filtered["filtered_row_count"]
    assert set((settings.dataset_storage_path / dataset).iterdir()) == files_before
    assert client.delete(f"/api/v1/datasets/{dataset}").status_code == 204
    assert client.post(f"/api/v1/datasets/{dataset}/export", json={}).status_code == 404


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_empty_export_valid_file_and_optional_audit(app_client, format):
    client, _ = app_client
    dataset = own_dataset(client)
    ready(client, dataset)
    response = client.post(f"/api/v1/datasets/{dataset}/export", json={"format": format, "scope": "filtered_data", "filters": [{"field": "channel", "op": "in", "values": ["Inexistente"]}], "include_transformations": False})
    if format == "csv":
        assert len(parse_csv(response)) == 1
    else:
        assert response.status_code == 200, response.text
        workbook = load_workbook(io.BytesIO(response.content))
        assert workbook.sheetnames == ["Datos", "Resumen"]
        assert len(list(workbook["Datos"].values)) == 1
        assert dict(list(workbook["Resumen"].values)[1:])["Filas exportadas"] == 0
    ready(client, dataset, [{"id": "drop_rows", "description": "Borrar selección", "params": {"row_ids": [1, 2, 3, 4]}, "destructive": True}])
    response = client.post(f"/api/v1/datasets/{dataset}/export", json={"format": "csv", "include_original_columns": True})
    assert len(parse_csv(response)) == 1


def test_export_errors_expiration_and_exporter_failure(app_client, monkeypatch):
    client, settings = app_client
    assert client.post("/api/v1/datasets/demo", json={"demo_id": "../invalid"}).json()["error"]["code"] == "DEMO_NOT_FOUND"
    dataset = own_dataset(client)
    endpoint = f"/api/v1/datasets/{dataset}/export"
    assert client.post(endpoint, json={}).json()["error"]["code"] == "STAGE_NOT_READY"
    ready(client, dataset)
    assert client.post(endpoint, json={"format": "pdf"}).json()["error"]["code"] == "EXPORT_FORMAT_UNSUPPORTED"
    assert client.post(endpoint, json={"filters": [{"field": "_row_id", "op": "in", "values": [1]}]}).json()["error"]["code"] == "FILTER_FIELD_INVALID"
    assert client.post(endpoint, json={"filters": [{"field": "amount", "op": "between", "values": []}]}).json()["error"]["code"] == "REQUEST_INVALID"
    assert client.post(endpoint, json={"scope": "filtered_data", "filters": [{"field": "amount", "op": "gte", "values": ["text"]}]}).json()["error"]["code"] == "EXPORT_FILTER_INVALID"
    from analytics_core.exports.xlsx import XLSXExporter
    def fail(*args):
        raise ValueError("C:/private/path secret cell")
    monkeypatch.setattr(XLSXExporter, "export", fail)
    failure = client.post(endpoint, json={"format": "xlsx"})
    assert failure.status_code == 500 and failure.json()["error"]["code"] == "EXPORT_FAILED"
    assert "private" not in failure.text and "secret" not in failure.text
    store = DatasetSessionStore(settings.dataset_storage_path, settings.dataset_ttl_minutes)
    session = store.get(dataset)
    session.expires_at = datetime.now(UTC) - timedelta(seconds=1)
    store.save(session)
    assert client.post(endpoint, json={}).json()["error"]["code"] == "DATASET_EXPIRED"
    assert not (settings.dataset_storage_path / dataset).exists()


def test_export_locale_and_cors(tmp_path):
    settings = Settings(dataset_storage_path=tmp_path, default_locale="en-US")
    with TestClient(create_app(settings)) as client:
        dataset = own_dataset(client)
        ready(client, dataset)
        rows = parse_csv(client.post(f"/api/v1/datasets/{dataset}/export", json={}), delimiter=",")
        assert rows[1][2] == "100.5"
        cors = client.options(f"/api/v1/datasets/{dataset}/cleaning", headers={"Origin": settings.frontend_origin, "Access-Control-Request-Method": "PUT"})
        assert cors.status_code == 200 and "PUT" in cors.headers["access-control-allow-methods"]
        download = client.post(f"/api/v1/datasets/{dataset}/export", json={}, headers={"Origin": settings.frontend_origin})
        assert "Content-Disposition" in download.headers["access-control-expose-headers"]
