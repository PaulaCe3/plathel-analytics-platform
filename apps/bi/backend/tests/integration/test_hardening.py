"""Small adversarial fixtures and complete-data/resource-limit regressions."""
import io
import json
import logging
import os
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime, timedelta
from threading import Lock
from zipfile import ZipFile, ZIP_DEFLATED
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.xml import DEFUSEDXML
from analytics_core.errors import AppError
from analytics_core.logging import JsonFormatter
from analytics_core.runtime import RuntimeGuard
from analytics_core.settings import Settings
from analytics_core.sessions.store import DatasetSessionStore
from bi.api.main import create_app


def settings(tmp_path, **changes):
    return Settings(dataset_storage_path=tmp_path, **changes)

def upload(client, data=b"fecha,importe\n2026-01-01,100\n", name="datos.csv", **kwargs):
    return client.post("/api/v1/datasets", files={"file": (name, data, "application/octet-stream")}, **kwargs)

def archive(entries):
    buffer = io.BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as out:
        for name, value in entries.items(): out.writestr(name, value)
    return buffer.getvalue()

@pytest.mark.parametrize("name,data,code", [
    ("datos.xls", b"text", "FILE_TYPE_UNSUPPORTED"),
    ("../datos.csv", b"a,b\n1,2", "FILE_TYPE_UNSUPPORTED"),
    ("datos.csv", b"MZexecutable", "FILE_CORRUPT"),
    ("datos.csv", b"PKzip", "FILE_CORRUPT"),
    ("datos.xlsx", b"not a zip", "FILE_CORRUPT"),
    ("datos.csv", b"a,b\n" + b"1,2\n" * 40000 + b"\x00", "FILE_CORRUPT"),
], ids=["extension", "filename", "executable", "zip-disguised", "false-xlsx", "late-nul"])
def test_rejects_invalid_uploads_without_session(tmp_path, name, data, code):
    with TestClient(create_app(settings(tmp_path))) as client:
        response = upload(client, data, name)
        assert response.json()["error"]["code"] == code
        assert not list(tmp_path.iterdir())

@pytest.mark.parametrize("extra,code", [
    ({"xl/vbaProject.bin": "macro"}, "FILE_TYPE_UNSUPPORTED"),
    ({"xl/externalLinks/externalLink1.xml": "external"}, "FILE_TYPE_UNSUPPORTED"),
    ({"xl/embeddings/oleObject1.bin": "ole"}, "FILE_TYPE_UNSUPPORTED"),
    ({"../escape.xml": "traversal"}, "FILE_CORRUPT"),
    ({"xl/large.xml": "x" * 30000}, "FILE_TOO_LARGE"),
])
def test_xlsx_adversarial(tmp_path, extra, code):
    data = archive({"[Content_Types].xml": "types", "xl/workbook.xml": "book", **extra})
    with TestClient(create_app(settings(tmp_path))) as client:
        assert upload(client, data, "datos.xlsx").json()["error"]["code"] == code
        assert not list(tmp_path.iterdir())

def test_zip_entry_and_expanded_limits(tmp_path):
    entries = {"[Content_Types].xml": "types", "xl/workbook.xml": "book", "third": "entry"}
    with TestClient(create_app(settings(tmp_path, max_zip_entries=2))) as client:
        assert upload(client, archive(entries), "datos.xlsx").status_code == 413
    from analytics_core.security.uploads import validate_xlsx
    path = tmp_path / "fixture.xlsx"; path.write_bytes(archive(entries))
    with pytest.raises(AppError) as caught: validate_xlsx(path, 5)
    assert caught.value.code == "FILE_TOO_LARGE"

@pytest.mark.parametrize("changes,data,code", [
    ({"max_rows": 1}, b"a,b\n1,2\n3,4\n", "TOO_MANY_ROWS"),
    ({"max_columns": 1}, b"a,b\n1,2\n", "TOO_MANY_COLUMNS"),
    ({"max_file_mb": 1}, b"a" * (1024 * 1024 + 1), "FILE_TOO_LARGE"),
], ids=["rows", "columns", "bytes"])
def test_real_file_limits_without_trusting_content_length(tmp_path, changes, data, code):
    with TestClient(create_app(settings(tmp_path, **changes))) as client:
        response = upload(client, data, headers={"Content-Length": "1"})
        assert response.json()["error"]["code"] == code
        assert not list(tmp_path.iterdir())

def test_content_length_early_rejection(tmp_path):
    with TestClient(create_app(settings(tmp_path, max_file_mb=1))) as client:
        assert upload(client, headers={"Content-Length": "9000000"}).status_code == 413
        assert upload(client, headers={"Content-Length": "wrong"}).status_code == 400

@pytest.mark.parametrize("dataset_id", ["../", "..\\", "a/b", "a\\b", "C:/absolute", "/absolute", "short", "%2e%2e%2f"])
def test_session_paths_stay_inside_root(tmp_path, dataset_id):
    store = DatasetSessionStore(tmp_path, 60)
    with pytest.raises(AppError): store.path(dataset_id, "raw.parquet")


def test_absolute_idle_expiration_and_permissions(tmp_path):
    store = DatasetSessionStore(tmp_path, 60, 240)
    session = store.create(); directory = tmp_path / session.dataset_id
    if os.name != "nt": assert directory.stat().st_mode & 0o777 == 0o700
    session.created_at = datetime.now(UTC) - timedelta(minutes=241)
    session.expires_at = datetime.now(UTC) + timedelta(minutes=60)
    store.save(session)
    with pytest.raises(AppError) as caught: store.get(session.dataset_id)
    assert caught.value.code == "DATASET_EXPIRED" and not directory.exists()
    session = store.create(); session.expires_at = datetime.now(UTC) - timedelta(seconds=1); store.save(session)
    assert store.purge_expired() == 1
    session = store.create(); session.created_at = datetime.now(UTC) - timedelta(minutes=239); store.save(session)
    touched = store.get(session.dataset_id)
    assert touched.expires_at <= touched.created_at + timedelta(minutes=240)


def test_startup_and_periodic_reaper(tmp_path):
    store = DatasetSessionStore(tmp_path, 60)
    first = store.create(); first.expires_at = datetime.now(UTC) - timedelta(seconds=1); store.save(first)
    with TestClient(create_app(settings(tmp_path, reaper_interval_seconds=0.05))) as client:
        assert not (tmp_path / first.dataset_id).exists()
        second = store.create(); second.expires_at = datetime.now(UTC) - timedelta(seconds=1); store.save(second)
        deadline = time.monotonic() + 3
        while (tmp_path / second.dataset_id).exists() and time.monotonic() < deadline: time.sleep(0.02)
        assert not (tmp_path / second.dataset_id).exists()
        assert client.get("/api/v1/health").json() == {"status": "ok"}


def test_rate_window_and_untrusted_forwarded_header(tmp_path):
    now = [0.0]; guard = RuntimeGuard(settings(tmp_path, uploads_per_hour=1), clock=lambda: now[0])
    guard.check_rate("one")
    with pytest.raises(AppError): guard.check_rate("one")
    now[0] = 3600; guard.check_rate("one")
    with TestClient(create_app(settings(tmp_path, uploads_per_hour=1))) as client:
        assert upload(client, headers={"X-Forwarded-For": "1.2.3.4"}).status_code == 201
        limited = client.post("/api/v1/datasets/demo", json={"demo_id": "retail_demo"}, headers={"X-Forwarded-For": "5.6.7.8"})
        assert limited.status_code == 429 and limited.json()["error"]["code"] == "RATE_LIMITED"
        assert client.get("/api/v1/health").status_code == 200

@pytest.mark.parametrize("limit", ["max_active_sessions", "max_sessions_per_ip"])
def test_active_session_capacity_freed_after_delete(tmp_path, limit):
    with TestClient(create_app(settings(tmp_path, **{limit: 1}))) as client:
        first = upload(client).json()["dataset_id"]
        assert upload(client).json()["error"]["code"] == "SESSION_LIMIT_REACHED"
        assert client.delete(f"/api/v1/datasets/{first}").status_code == 204
        assert client.get(f"/api/v1/datasets/{first}").status_code == 404
        assert upload(client).status_code == 201

@pytest.mark.parametrize("method", ["POST", "PUT", "DELETE"])
def test_restricted_cors_preflight(tmp_path, method):
    with TestClient(create_app(settings(tmp_path))) as client:
        headers = {"Origin": "http://localhost:3000", "Access-Control-Request-Method": method, "Access-Control-Request-Headers": "Content-Type,X-Request-ID"}
        allowed = client.options("/api/v1/datasets", headers=headers)
        assert allowed.status_code == 200 and allowed.headers["access-control-allow-origin"] == headers["Origin"]
        assert "access-control-allow-credentials" not in allowed.headers
        headers["Origin"] = "https://untrusted.example"
        assert client.options("/api/v1/datasets", headers=headers).status_code == 400


def test_sanitized_unknown_error_request_id_and_logs(tmp_path):
    app = create_app(settings(tmp_path))
    @app.get("/failure")
    def fail(): raise ValueError("private customer SECRET_CELL C:/private/path")
    with TestClient(app, raise_server_exceptions=False) as client:
        response = client.get("/failure", headers={"X-Request-ID": "r" * 500})
        assert response.status_code == 500
        assert response.json()["error"]["code"] == "INTERNAL_ERROR"
        assert response.json()["error"]["request_id"] == response.headers["x-request-id"]
        assert "SECRET_CELL" not in response.text and "traceback" not in response.text
        assert response.headers["x-content-type-options"] == "nosniff"
    try: raise ValueError("SECRET_CELL")
    except ValueError:
        import sys
        record = logging.LogRecord("test", logging.ERROR, "private", 1, "safe error", (), sys.exc_info())
    output = JsonFormatter().format(record)
    assert "SECRET_CELL" not in output and "ValueError" in output


def test_heavy_concurrency_bound(tmp_path):
    guard = RuntimeGuard(settings(tmp_path, heavy_concurrency=2)); lock = Lock(); current = 0; peak = 0
    def work(_):
        nonlocal current, peak
        with guard.heavy():
            with lock: current += 1; peak = max(peak, current)
            time.sleep(0.02)
            with lock: current -= 1
    with ThreadPoolExecutor(max_workers=6) as pool: list(pool.map(work, range(6)))
    assert peak == 2


def test_xlsx_blank_rows_limits_and_formulas_never_executed(tmp_path):
    assert DEFUSEDXML
    workbook = Workbook(); sheet = workbook.active
    sheet.append(["fecha", "importe"]); sheet.append(["2026-01-01", 100]); sheet.append([None, None]); sheet.append(["2026-01-02", "=1+1"])
    buffer = io.BytesIO(); workbook.save(buffer)
    with TestClient(create_app(settings(tmp_path, max_rows=1))) as client:
        assert upload(client, buffer.getvalue(), "datos.xlsx").json()["error"]["code"] == "TOO_MANY_ROWS"
    with TestClient(create_app(settings(tmp_path))) as client:
        response = upload(client, buffer.getvalue(), "datos.xlsx"); assert response.status_code == 201
        dataset = response.json()["dataset_id"]
        preview = client.get(f"/api/v1/datasets/{dataset}/preview").json()
        assert preview["rows"][1]["c02"] is None


def test_full_dataset_metrics_performance_smoke(tmp_path):
    data = "date,concept,amount,channel,currency\n" + "".join(f"2026-01-01,Product {index},1,Online,ARS\n" for index in range(2000))
    started = time.monotonic()
    with TestClient(create_app(settings(tmp_path, profile_sample_rows=10))) as client:
        response = upload(client, data.encode()); assert response.status_code == 201
        dataset = response.json()["dataset_id"]
        assert len(client.get(f"/api/v1/datasets/{dataset}/preview").json()["rows"]) == 50
        mappings = [{"column_key": f"c{i:02}", "target_field": field, "disposition": "canonical"} for i, field in enumerate(["date", "concept", "amount", "channel", "currency"], 1)]
        assert client.put(f"/api/v1/datasets/{dataset}/mapping", json={"profile_id": "retail_ecommerce", "mappings": mappings}).status_code == 200
        assert client.post(f"/api/v1/datasets/{dataset}/validate").status_code == 200
        assert client.put(f"/api/v1/datasets/{dataset}/cleaning", json={"actions": []}).status_code == 200
        dashboard = client.post(f"/api/v1/datasets/{dataset}/dashboard", json={}).json()
        assert dashboard["data"]["kpi_revenue"]["value"] == 2000
        assert dashboard["filtered_row_count"] == 2000
        assert sum(point[1] for point in dashboard["data"]["revenue_by_concept"]["series"][0]["points"]) == 2000
        assert time.monotonic() - started < 60


def test_csv_wide_data_row_does_not_become_a_silent_index(tmp_path):
    with TestClient(create_app(settings(tmp_path))) as client:
        response = upload(client, b"a,b\n1,2,3\n")
        assert response.status_code == 422 and response.json()["error"]["code"] == "FILE_CORRUPT"


def test_xlsx_source_removed_after_mapping_and_all_artifacts_deleted(tmp_path):
    workbook = Workbook(); sheet = workbook.active
    sheet.append(["date", "amount"]); sheet.append(["2026-01-01", 100])
    buffer = io.BytesIO(); workbook.save(buffer)
    with TestClient(create_app(settings(tmp_path))) as client:
        response = upload(client, buffer.getvalue(), "datos.xlsx"); assert response.status_code == 201
        dataset = response.json()["dataset_id"]; directory = tmp_path / dataset
        assert (directory / "source.bin").exists()
        assert client.patch(f"/api/v1/datasets/{dataset}", json={"sheet": sheet.title}).status_code == 200
        mappings = [{"column_key": "c01", "target_field": "date", "disposition": "canonical"}, {"column_key": "c02", "target_field": "amount", "disposition": "canonical"}]
        assert client.put(f"/api/v1/datasets/{dataset}/mapping", json={"profile_id": "custom", "mappings": mappings}).status_code == 200
        assert not (directory / "source.bin").exists() and (directory / "raw.parquet").exists()
        assert client.patch(f"/api/v1/datasets/{dataset}", json={"sheet": sheet.title}).status_code == 409
        assert client.post(f"/api/v1/datasets/{dataset}/validate").status_code == 200
        assert (directory / "canonical.parquet").exists()
        assert client.delete(f"/api/v1/datasets/{dataset}").status_code == 204
        assert not directory.exists()


def test_reaper_does_not_delete_active_session(tmp_path):
    store = DatasetSessionStore(tmp_path, 60); session = store.create()
    session.expires_at = datetime.now(UTC) - timedelta(seconds=1); store.save(session)
    with store.lock(session.dataset_id):
        with ThreadPoolExecutor(max_workers=1) as pool: assert pool.submit(store.purge_expired).result() == 0
    assert store.purge_expired() == 1


def test_proxy_headers_only_from_explicit_peer_and_no_ip_persistence(tmp_path):
    app = create_app(settings(tmp_path, trusted_proxy_ips=("testclient",), uploads_per_hour=1))
    with TestClient(app) as client:
        for ip in ["1.2.3.4", "5.6.7.8"]:
            response = upload(client, headers={"X-Forwarded-For": ip}); assert response.status_code == 201
            metadata = (tmp_path / response.json()["dataset_id"] / "session.json").read_text()
            assert ip not in metadata
        assert upload(client, headers={"X-Forwarded-For": "1.2.3.4"}).status_code == 429


def test_parse_error_is_generic_and_logs_exclude_values(tmp_path, caplog):
    app = create_app(settings(tmp_path))
    logger = logging.getLogger("data_analytics_platform"); logger.addHandler(caplog.handler)
    with TestClient(app) as client:
        logger.addHandler(caplog.handler)
        response = upload(client, b"a,b\nPRIVATE,1,2\n", "PERSONAL_FILENAME.csv")
        assert response.status_code == 422
    logger.removeHandler(caplog.handler)
    assert caplog.records
    for record in caplog.records:
        output = JsonFormatter().format(record)
        assert "PRIVATE" not in output and "PERSONAL_FILENAME" not in output


def test_complete_chunked_body_limit_including_non_file_parts(tmp_path):
    with TestClient(create_app(settings(tmp_path, max_file_mb=1))) as client:
        response = client.post("/api/v1/datasets", content=iter([b"x" * 700000, b"y" * 700000]), headers={"Content-Type": "multipart/form-data; boundary=x", "Content-Length": "1"})
        assert response.status_code == 413
        assert response.json()["error"]["code"] == "FILE_TOO_LARGE"
        assert response.headers["x-request-id"] == response.json()["error"]["request_id"]
        assert not list(tmp_path.iterdir())


def test_defused_xml_entity_fixture_rejected(tmp_path):
    workbook = Workbook(); workbook.active.append(["date", "amount"]); workbook.active.append(["2026-01-01", 1])
    buffer = io.BytesIO(); workbook.save(buffer)
    with ZipFile(io.BytesIO(buffer.getvalue())) as source:
        entries = {name: source.read(name) for name in source.namelist()}
    entries["xl/workbook.xml"] = b'<!DOCTYPE workbook [<!ENTITY xxe SYSTEM "file:///nonexistent-fixture">]>' + entries["xl/workbook.xml"]
    with TestClient(create_app(settings(tmp_path))) as client:
        response = upload(client, archive(entries), "entity.xlsx")
        assert response.status_code == 422 and response.json()["error"]["code"] == "FILE_CORRUPT"
