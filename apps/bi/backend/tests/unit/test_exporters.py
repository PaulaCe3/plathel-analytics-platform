from datetime import date
from io import BytesIO
import pytest
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook
from analytics_core.cleaning.models import TransformationLog
from analytics_core.exports import ExportOptions, ExportSource, ExporterRegistry
from analytics_core.exports.csv import CSVExporter
from analytics_core.exports.xlsx import XLSXExporter
from analytics_core.exports.security import safe_cell, unique_headers
from analytics_core.engine.pandas_impl import PandasDataEngine
from analytics_core.engine.query import FilterClause, MeasureSpec, QuerySpec


@pytest.mark.parametrize("value", ["=cmd", "+cmd", "-cmd", "@cmd", "\tcmd", "\rcmd"])
def test_text_sanitizer(value):
    assert safe_cell(value) == "'" + value
    assert safe_cell(-42.5) == -42.5
    assert safe_cell(-42) == -42
    assert safe_cell(None) is None


def test_headers_collision_deterministic():
    assert unique_headers(["Importe", "Importe", "Importe (2)", "importe", ""]) == ("Importe", "Importe (2)", "Importe (2) (2)", "importe (3)", "Columna")


def test_csv_incremental_stream_and_locale():
    reads = []
    def rows():
        for i in range(3):
            reads.append(i)
            yield [date(2026, 1, 1), "Córdoba", -42.5, "=1+1"]
    source = ExportSource(("Fecha", "Ubicación", "Valor", "Texto"), rows, {}, TransformationLog())
    iterator = CSVExporter().export(source, ExportOptions())
    assert next(iterator).startswith(b"\xef\xbb\xbf") and not reads
    first = next(iterator)
    assert reads == [0] and "-42,5" in first.decode("utf-8") and "'=1+1" in first.decode("utf-8")
    assert len(list(iterator)) == 2


def test_xlsx_security_in_headers_summary_and_audit():
    from analytics_core.cleaning.models import Transformation
    log = TransformationLog(transformations=[Transformation(id="x", seq=1, action_id="trim_whitespace", kind="normalize", rows_before=1, rows_after=1, rows_affected=1, summary="=cmd", automatic=False)])
    source = ExportSource(("@header",), lambda: [["+cmd"]], {"=key": "-cmd", "Número": -42.5}, log)
    binary = b"".join(XLSXExporter().export(source, ExportOptions(format="xlsx")))
    workbook = load_workbook(BytesIO(binary), data_only=False)
    assert workbook["Datos"]["A1"].value == "'@header"
    assert workbook["Datos"]["A2"].value == "'+cmd"
    assert workbook["Resumen"]["A2"].value == "'=key" and workbook["Resumen"]["B2"].value == "'-cmd"
    assert workbook["Resumen"]["B3"].value == -42.5
    assert workbook["Registro de cambios"]["J2"].value == "'=cmd"
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)


@pytest.mark.parametrize("clause", [FilterClause(field="channel", op="in", values=["Online"]), FilterClause(field="channel", op="not_in", values=["Local"]), FilterClause(field="amount", op="between", values=[2, 30]), FilterClause(field="amount", op="gte", values=[12]), FilterClause(field="amount", op="lte", values=[12]), FilterClause(field="channel", op="contains", values=["on"])])
def test_row_projection_matches_query_for_every_filter_and_full_dataset(tmp_path, clause):
    data = {"_row_id": list(range(1, 9001)), "channel": ["Online" if i % 2 else "Local" for i in range(9000)], "amount": [float(i) for i in range(9000)]}
    path = tmp_path / "canonical.parquet"
    pq.write_table(pa.table(data), path)
    engine = PandasDataEngine()
    rows = list(engine.iter_rows(path, QuerySpec(measures=[], filters=[clause]), ["_row_id", "amount"]))
    query = engine.run_query(path, QuerySpec(measures=[MeasureSpec(alias="rows", aggregation="count"), MeasureSpec(alias="total", aggregation="sum", field="amount")], filters=[clause]))
    assert len(rows) == query.rows[0]["rows"]
    assert sum(row["amount"] for row in rows) == query.rows[0]["total"]
    assert all(set(row) == {"_row_id", "amount"} for row in rows)


def test_registry_extends_without_changing_dispatch():
    class AdditionalExporter:
        format = "test_format"
        content_type = "application/octet-stream"
        def export(self, source, options):
            return iter([b"test"])
    registry = ExporterRegistry()
    registry.register(AdditionalExporter())
    assert b"".join(registry.get("test_format").export(None, None)) == b"test"
    with pytest.raises(ValueError):
        registry.register(AdditionalExporter())


@pytest.mark.parametrize("format", ["csv", "xlsx"])
@pytest.mark.parametrize("text", ["\rformula", "\tformula", "\nformula"])
def test_control_character_text_protected_in_each_exporter(format, text):
    source = ExportSource(("Texto", "Número"), lambda: [[text, -42.5]], {}, TransformationLog())
    exporter = CSVExporter() if format == "csv" else XLSXExporter()
    binary = b"".join(exporter.export(source, ExportOptions(format=format)))
    if format == "csv":
        import csv
        from io import StringIO
        rows = list(csv.reader(StringIO(binary.decode("utf-8-sig")), delimiter=";"))
        assert rows[1][0] == "'" + text and rows[1][1] == "-42,5"
    else:
        workbook = load_workbook(BytesIO(binary))
        assert workbook["Datos"]["A2"].value.startswith("'")
        assert workbook["Datos"]["B2"].value == -42.5


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_header_collisions_after_formula_neutralization(format):
    source = ExportSource(("=Título", "'=Título"), lambda: [[1, 2]], {}, TransformationLog())
    exporter = CSVExporter() if format == "csv" else XLSXExporter()
    binary = b"".join(exporter.export(source, ExportOptions(format=format)))
    if format == "csv":
        import csv
        from io import StringIO
        headers = next(csv.reader(StringIO(binary.decode("utf-8-sig")), delimiter=";"))
    else:
        headers = next(load_workbook(BytesIO(binary))["Datos"].values)
    assert list(headers) == ["'=Título", "'=Título (2)"]
